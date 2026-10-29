"""态势感知(安全感知平台)导出表解析：支持 CSV / XLSX，自动识别关键列。"""
import csv
import os
from . import timeutil

# 字段 -> 可能的表头别名(已做小写+去空格归一化)
HEADER_MAP = {
    "time": ["时间", "发现时间", "告警时间", "检测时间", "事件时间", "time", "时间戳", "timestamp", "日期", "发生时间"],
    "src_ip": ["源ip", "源ip地址", "源地址", "src_ip", "srcip", "源终端ip", "终端ip", "源"],
    "dst_ip": ["目的ip", "目的ip地址", "目的地址", "dst_ip", "dstip", "目的终端ip", "目的"],
    "alert": ["告警名称", "告警", "名称", "描述", "alert", "告警类型", "事件名称"],
    "threat": ["威胁类型", "威胁", "threat", "type", "攻击类型"],
    "result": ["攻击结果", "结果", "result", "处置结果"],
    "src_port": ["源端口", "src_port", "srcport"],
    "dst_port": ["目的端口", "dst_port", "dstport"],
    "status": ["处置状态", "状态", "status", "disposition"],
    "host": ["主机名", "hostname", "host"],
    "app": ["应用名称", "应用", "appname", "application"],
}

# 深澜导出表头别名（在线表 / 上网明细 / 认证日志 CSV、XLSX）
SHENLAN_MAP = {
    "account": ["账号", "account", "user", "username", "userid", "用户账号", "用户名", "上网账号", "客户端账号"],
    "name": ["用户姓名", "姓名", "name", "realname", "real_name", "真实姓名"],
    "ip": ["ip地址", "ip", "ipaddr", "ip_address", "上网ip", "客户端ip", "终端ip", "上线ip"],
    "mac": ["mac地址", "mac", "macaddr", "mac_address", "终端mac"],
    "online_time": ["上线时间", "在线时间", "online_time", "logintime", "starttime", "start_time", "acctstarttime", "acctstart", "login_time", "登录时间", "认证时间", "开始时间"],
    "offline_time": ["下线时间", "离线时间", "offline_time", "stoptime", "endtime", "end_time", "acctstoptime", "acctstop", "结束时间", "停止时间", "断开时间"],
}


def _norm_header(h):
    return (h or "").strip().lower().replace(" ", "").replace("（", "(").replace("）", ")")


# 收集所有已知表头别名，用于自动定位真实表头所在行（跳过"日志检索高级模式"等元信息行）
KNOWN_HEADERS = set()
for _m in (HEADER_MAP, SHENLAN_MAP):
    for _variants in _m.values():
        for _v in _variants:
            KNOWN_HEADERS.add(_norm_header(_v))
# 额外补充一些醒目的表头标记
for _extra in ("序号", "时间", "源ip", "目的ip", "描述", "攻击类型", "动作", "状态码",
               "数据来源", "主机名", "应用名称", "数据包", "账号", "ip地址", "用户名"):
    KNOWN_HEADERS.add(_norm_header(_extra))


def detect_columns(headers):
    norm = {_norm_header(h): h for h in headers}
    mapping = {}
    for field, variants in HEADER_MAP.items():
        for v in variants:
            if v in norm:
                mapping[field] = norm[v]
                break
    return mapping


def _as_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _score_header(row):
    """给某一行打分：命中已知表头别名越多分数越高；含 序号+时间+源/目的IP 组合额外加分。"""
    norm = [_norm_header(c) for c in row if c is not None]
    score = sum(1 for n in norm if n in KNOWN_HEADERS)
    text = " ".join(norm)
    if "时间" in text and ("源ip" in text or "目的ip" in text or "ip" in text):
        score += 3
    if "序号" in text:
        score += 1
    return score


def _find_header_index(rows):
    """在若干行里找出真正的表头行（跳过'日志检索高级模式/过滤条件'等元信息行）。"""
    best, best_score = -1, 2  # 阈值：至少要有一定命中
    for i, r in enumerate(rows):
        s = _score_header(r)
        if s > best_score:
            best, best_score = i, s
    return best


def _read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    else:
        # CSV / TSV
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            delim = "\t" if sample.count("\t") > sample.count(",") else ","
            reader = csv.reader(f, delimiter=delim)
            rows = list(reader)
    if not rows:
        return [], []

    hi = _find_header_index(rows)
    if hi < 0:  # 没找到明显表头，退化为首行
        hi = 0
    headers = [_as_str(c) for c in rows[hi]]
    data = []
    for r in rows[hi + 1:]:
        if ext in (".xlsx", ".xls"):
            if all(c is None for c in r):
                continue
            data.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
        else:
            if all((c or "").strip() == "" for c in r):
                continue
            data.append({headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers))})
    return headers, data


def parse_situational_file(path):
    headers, data = _read_rows(path)
    colmap = detect_columns(headers)
    out = []
    for row in data:
        rec = {}
        for field, h in colmap.items():
            rec[field] = _as_str(row.get(h))
        rec["_raw"] = {
            h: _as_str(row.get(h))
            for h in headers
            if not isinstance(row.get(h), (list, dict))
        }
        out.append(rec)
    return {"columns": colmap, "records": out, "total": len(out), "headers": headers}


def parse_shenlan_file(path):
    """解析深澜导出表（在线表/认证日志），输出标准化记录。"""
    headers, data = _read_rows(path)
    colmap = detect_columns_shenlan(headers)
    out = []
    for row in data:
        rec = {}
        for field, h in colmap.items():
            rec[field] = _as_str(row.get(h))
        out.append(rec)
    return {"columns": colmap, "records": out, "total": len(out)}


def detect_columns_shenlan(headers):
    norm = {_norm_header(h): h for h in headers}
    mapping = {}
    for field, variants in SHENLAN_MAP.items():
        for v in variants:
            if v in norm:
                mapping[field] = norm[v]
                break
    return mapping
